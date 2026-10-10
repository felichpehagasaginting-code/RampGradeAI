import io
from datetime import datetime
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class FactoryExcelExporter:
    """
    Generator Laporan Harian Pengolahan (LHP) & Rekap Mutu TBS Pabrik Kelapa Sawit (.xlsx)
    Format resmi untuk manajemen pabrik, mandor timbang, dan departemen akuntansi.
    """

    @classmethod
    def generate_lhp_excel(cls, sessions: list) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LHP Rekap Mutu TBS"

        # Palet Warna PKS
        primary_fill = PatternFill(start_color="064E3B", end_color="064E3B", fill_type="solid")  # Hijau Tua Sawit
        header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")   # Teal Tua
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")    # Abu Lembut
        summary_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")  # Hijau Pastel
        
        white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)
        number_font = Font(name="Consolas", size=10)

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # 1. Judul & Kop Laporan
        ws.merge_cells("A1:M1")
        title_cell = ws["A1"]
        title_cell.value = "LAPORAN REKAPITULASI SORTASI TBS & ESTIMASI OER PKS"
        title_cell.font = title_font
        title_cell.fill = primary_fill
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 32

        ws.merge_cells("A2:M2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Sistem SMART GRADE • Politeknik Citra Widya Edukasi (PCWE) • Dicetak: {datetime.now().strftime('%d-%m-%Y %H:%M:%S WIB')}"
        sub_cell.font = Font(name="Calibri", size=9, italic=True, color="475569")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        ws.append([]) # Baris kosong 3

        # 2. Header Tabel
        headers = [
            "No", "Kode Sesi Grading", "No. SPB", "Plat Truk", "Nama Pekebun",
            "Tonase (Kg)", "Total TBS", "Matang (%)", "Kurang (%)", "Mentah (%)",
            "Lewat (%)", "OER (%)", "Nilai Payout (Rp)"
        ]
        ws.append(headers)
        ws.row_dimensions[4].height = 26

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.font = white_bold
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # 3. Baris Data Transaksi
        total_tonnage = 0.0
        total_bunches = 0
        total_payout = 0.0

        for idx, s in enumerate(sessions, 1):
            row_num = 4 + idx
            order = s.delivery_order
            farmer = order.farmer if order else None
            oer_obj = s.oer_assessment
            
            tonnage = order.net_weight_kg if order else 0.0
            payout = oer_obj.total_payout if oer_obj else 0.0
            
            total_tonnage += tonnage
            total_bunches += s.total_bunches
            total_payout += payout

            # Persentase fraksi
            tot = s.total_bunches or 1
            p_ripe = (s.count_ripe / tot) * 100.0
            p_under = (s.count_underripe / tot) * 100.0
            p_unripe = (s.count_unripe / tot) * 100.0
            p_over = (s.count_overripe / tot) * 100.0
            oer_val = oer_obj.estimated_oer_percentage if oer_obj else 21.5

            row_data = [
                idx,
                s.session_code,
                order.spb_number if order else "-",
                order.truck_plate_number if order else "-",
                farmer.name if farmer else "Pekebun Swadaya",
                tonnage,
                s.total_bunches,
                round(p_ripe, 1),
                round(p_under, 1),
                round(p_unripe, 1),
                round(p_over, 1),
                round(oer_val, 2),
                round(payout, 0)
            ]
            ws.append(row_data)
            ws.row_dimensions[row_num].height = 20

            # Formatting
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if idx % 2 == 0:
                    cell.fill = zebra_fill
                
                # Alignment khusus
                if col_idx in [1, 3, 4]:
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx in [6, 7, 8, 9, 10, 11, 12, 13]:
                    cell.alignment = Alignment(horizontal="right")
                    cell.font = number_font
                    if col_idx in [6, 7]:
                        cell.number_format = "#,##0"
                    elif col_idx in [8, 9, 10, 11, 12]:
                        cell.number_format = "0.0"
                    elif col_idx == 13:
                        cell.number_format = "Rp #,##0"

        # 4. Baris Total / Summary
        summary_row = 5 + len(sessions)
        ws.append([
            "TOTAL / RATA-RATA", "", "", "", "",
            total_tonnage, total_bunches, "", "", "", "",
            round(sum(s.oer_assessment.estimated_oer_percentage for s in sessions if s.oer_assessment) / max(1, len(sessions)), 2),
            total_payout
        ])
        ws.row_dimensions[summary_row].height = 24

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=summary_row, column=col_idx)
            cell.font = bold_font
            cell.fill = summary_fill
            cell.border = thin_border
            if col_idx in [6, 7, 12, 13]:
                cell.alignment = Alignment(horizontal="right")
                if col_idx == 13:
                    cell.number_format = "Rp #,##0"

        # Auto-fit Column Widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(11, min(max_len + 3, 30))

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
